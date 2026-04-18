"""
Unit tests for Ariadne ingestors.

Run with: python -m pytest tests/ -v
"""

import pytest
import tempfile
import os
from pathlib import Path

from ariadne.ingest import (
    BaseIngestor,
    Document,
    MarkdownIngestor,
    TxtIngestor,
    ExcelIngestor,
    CsvIngestor,
    SourceType,
)


class TestDocument:
    """Test the Document data model."""
    
    def test_document_creation(self):
        """Document can be created with required fields."""
        doc = Document(content="Hello, world!")
        assert doc.content == "Hello, world!"
        assert doc.source_type == SourceType.UNKNOWN
        assert doc.chunk_index == 0
        assert doc.total_chunks == 1
    
    def test_document_with_metadata(self):
        """Document accepts custom metadata."""
        doc = Document(
            content="Test content",
            source_type=SourceType.MARKDOWN,
            source_path="/path/to/file.md",
            metadata={"author": "Test User", "title": "Test Doc"}
        )
        assert doc.metadata["author"] == "Test User"
        assert doc.metadata["title"] == "Test Doc"
    
    def test_document_doc_id(self):
        """Document generates stable doc_id."""
        doc = Document(
            content="Same content",
            source_path="/path/to/file.md",
            chunk_index=0
        )
        doc2 = Document(
            content="Same content",
            source_path="/path/to/file.md",
            chunk_index=0
        )
        # Same path + index should produce same ID
        assert doc.doc_id == doc2.doc_id
    
    def test_document_empty_content_raises(self):
        """Document rejects empty content."""
        with pytest.raises(ValueError):
            Document(content="")
        with pytest.raises(ValueError):
            Document(content="   ")


class TestMarkdownIngestor:
    """Test Markdown ingestion."""
    
    @pytest.fixture
    def temp_md_file(self):
        """Create a temporary markdown file."""
        with tempfile.NamedTemporaryFile(
            mode='w', suffix='.md', delete=False, encoding='utf-8'
        ) as f:
            f.write("# Hello\n\nThis is a test.\n\n## Section\n\nMore content.")
            path = f.name
        yield path
        os.unlink(path)
    
    def test_ingest_markdown(self, temp_md_file):
        """Markdown file is parsed into documents."""
        ingestor = MarkdownIngestor()
        docs = ingestor.ingest(temp_md_file)
        assert len(docs) >= 1
        assert any("Hello" in doc.content for doc in docs)


class TestTxtIngestor:
    """Test plain text ingestion."""
    
    @pytest.fixture
    def temp_txt_file(self):
        """Create a temporary text file."""
        with tempfile.NamedTemporaryFile(
            mode='w', suffix='.txt', delete=False, encoding='utf-8'
        ) as f:
            f.write("First paragraph.\n\nSecond paragraph.\n\nThird paragraph.")
            path = f.name
        yield path
        os.unlink(path)
    
    def test_ingest_txt(self, temp_txt_file):
        """Text file is parsed into documents."""
        ingestor = TxtIngestor()
        docs = ingestor.ingest(temp_txt_file)
        assert len(docs) >= 1
        assert all(isinstance(doc, Document) for doc in docs)


class TestExcelIngestor:
    """Test Excel spreadsheet ingestion."""
    
    @pytest.fixture
    def temp_xlsx_file(self):
        """Create a temporary Excel file using openpyxl."""
        try:
            import openpyxl
        except ImportError:
            pytest.skip("openpyxl not installed")
        
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Test Sheet"
        
        # Add headers
        ws.append(["Name", "Age", "City"])
        # Add data rows
        ws.append(["Alice", 30, "Beijing"])
        ws.append(["Bob", 25, "Shanghai"])
        ws.append(["Charlie", 35, "Shenzhen"])
        
        # Create temp file
        fd, path = tempfile.mkstemp(suffix='.xlsx')
        os.close(fd)
        wb.save(path)
        wb.close()
        
        yield path
        os.unlink(path)
    
    def test_ingest_excel(self, temp_xlsx_file):
        """Excel file is parsed into documents."""
        try:
            import openpyxl
        except ImportError:
            pytest.skip("openpyxl not installed")
        
        ingestor = ExcelIngestor()
        docs = ingestor.ingest(temp_xlsx_file)
        assert len(docs) >= 1
        assert all(isinstance(doc, Document) for doc in docs)
        # Check that source_type is correctly set
        assert docs[0].source_type == SourceType.EXCEL


class TestCsvIngestor:
    """Test CSV file ingestion."""
    
    @pytest.fixture
    def temp_csv_file(self):
        """Create a temporary CSV file."""
        fd, path = tempfile.mkstemp(suffix='.csv')
        with os.fdopen(fd, 'w', encoding='utf-8') as f:
            f.write("Name,Age,City\n")
            f.write("Alice,30,Beijing\n")
            f.write("Bob,25,Shanghai\n")
            f.write("Charlie,35,Shenzhen\n")
        yield path
        os.unlink(path)
    
    def test_ingest_csv(self, temp_csv_file):
        """CSV file is parsed into documents."""
        ingestor = CsvIngestor()
        docs = ingestor.ingest(temp_csv_file)
        assert len(docs) >= 1
        # First doc should contain headers
        header_doc = docs[0]
        assert "Name" in header_doc.content
        assert "Age" in header_doc.content
        assert "City" in header_doc.content
        # Check that source_type is correctly set
        assert docs[0].source_type == SourceType.CSV
    
    @pytest.fixture
    def temp_semicolon_csv(self):
        """Create a CSV file with semicolon delimiter."""
        fd, path = tempfile.mkstemp(suffix='.csv')
        with os.fdopen(fd, 'w', encoding='utf-8') as f:
            f.write("Name;Age;City\n")
            f.write("Alice;30;Beijing\n")
        yield path
        os.unlink(path)
    
    def test_delimiter_detection(self, temp_semicolon_csv):
        """CSV ingestor auto-detects semicolon delimiter."""
        ingestor = CsvIngestor()
        docs = ingestor.ingest(temp_semicolon_csv)
        # Should still parse correctly with auto-detection
        assert any("Name" in doc.content for doc in docs)


class TestBaseIngestor:
    """Test base ingestor functionality."""
    
    def test_file_not_found(self):
        """Ingestor raises FileNotFoundError for missing files."""
        ingestor = TxtIngestor()
        with pytest.raises(FileNotFoundError):
            ingestor.ingest("/nonexistent/path/file.txt")
    
    def test_chunk_text(self):
        """chunk_text helper splits long text correctly."""
        text = "A" * 1000
        chunks = BaseIngestor.chunk_text(text, max_chars=200, overlap=20)
        assert len(chunks) > 1
        # Check overlap
        assert chunks[0][-20:] == chunks[1][:20]
    
    def test_chunk_text_short(self):
        """chunk_text returns single chunk for short text."""
        text = "Short text"
        chunks = BaseIngestor.chunk_text(text, max_chars=500)
        assert len(chunks) == 1
        assert chunks[0] == text
